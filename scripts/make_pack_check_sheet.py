"""Build the pack-checking sheet: a stratified random sample of the real Open Food Facts products, with
what the system stored and concluded pre-filled next to blank columns for what the printed pack says.

    ./run.sh api &                                   # needs the seeded API
    python scripts/make_pack_check_sheet.py          # writes data/verified/pack_check/

Synthetic records are never sampled. The sample is fixed by SEED so the team and the report refer to
the same products. Nothing is filled in on the checker side: those columns come only from real packs.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import urllib.request
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "verified" / "pack_check"
SEED = 2026
# Main sample of 150 (at least 20 per category so per-category agreement can be reported) plus 10
# reserves per category, used in order when a sampled product cannot be found in shops.
QUOTA = {"breakfast_cereals": 20, "biscuits_cookies": 30, "instant_noodles": 20, "beverages": 50,
         "namkeen_chips": 30}
RESERVE = 10

NUTRIENTS = [("energy", "kcal"), ("protein", "g"), ("carbohydrate", "g"), ("total_sugars", "g"),
             ("added_sugars", "g"), ("total_fat", "g"), ("saturated_fat", "g"), ("trans_fat", "g"),
             ("fibre", "g"), ("sodium", "mg")]

ID_COLS = ["check_id", "sample", "reserve_order", "category", "product_id", "off_code", "off_url", "name",
           "brand", "pack_size_stored"]
SYS_COLS = (["sys_basis"] + [f"sys_{n}_{u}" for n, u in NUTRIENTS] +
            ["sys_ingredients", "sys_ingredient_count", "sys_ingredients_uncertain", "sys_claims",
             "sys_overall", "sys_grade", "sys_data_status", "sys_alt_1", "sys_alt_2", "sys_alt_3"])
CHECK_COLS = (["checked_by", "checked_on", "pack_seen_where", "pack_found", "category_correct", "label_pack_size",
               "label_basis"] +
              [f"label_{n}_{u}" for n, u in NUTRIENTS] +
              ["ingredients_match", "ingredients_notes", "label_claims", "claims_assessment",
               "rec_alt_1_quality", "rec_alt_2_quality", "rec_alt_3_quality", "rec_notes", "notes"])

CHOICES = {
    "pack_seen_where": ["shop", "own pack", "brand website photo", "retailer photo"],
    "pack_found": ["yes", "different variant", "not found"],
    "category_correct": ["yes", "no"],
    "label_basis": ["per_100g", "per_100ml", "per serving only"],
    "ingredients_match": ["same", "minor differences", "major differences", "not printed"],
    "claims_assessment": ["agree", "disagree", "missed claim", "no claims"],
    "rec_alt_1_quality": ["1", "2", "3", "4", "5"],
    "rec_alt_2_quality": ["1", "2", "3", "4", "5"],
    "rec_alt_3_quality": ["1", "2", "3", "4", "5"],
}

GUIDE = [
    ("What this is", "A fixed random sample of real Open Food Facts India products (seed %d). The sys_ columns show "
                     "what the system stored and concluded. Fill the blank columns from the printed pack only." % SEED),
    ("Which rows", "Check every 'main' row (150). If a product cannot be found, set pack_found = not found and use the "
                   "next unused 'reserve' row of the same category, in reserve_order."),
    ("Evidence", "Use the physical pack or a clear photo of it. Keep photos in the team drive, not in git (they may show "
                 "people or shops). Never copy a value from Open Food Facts or a retailer listing into a label_ column."),
    ("Category", "category_correct = no when the product is not really in the stored category (for example a "
                 "namkeen filed under breakfast cereals). Categories come from Open Food Facts tags and can be wrong."),
    ("Nutrition", "Copy values per 100 g or 100 ml exactly as printed, in the units of the column header (sodium in mg). "
                  "If only per-serving values are printed, set label_basis = per serving only and note the serving size. "
                  "Leave a cell empty if the value is not printed; do not write 0."),
    ("Ingredients", "Compare the printed list with sys_ingredients: same, minor differences (spelling, order of two "
                    "items, a missing percentage) or major differences (different recipe, missing items)."),
    ("Claims", "Write every front-of-pack claim in label_claims, separated by ';'. claims_assessment: agree if "
               "sys_claims captured them and the status looks right, disagree if a status looks wrong, missed claim "
               "if a printed claim is absent from sys_claims."),
    ("Recommendations", "Rate each suggested alternative 1-5: 5 = a sensible swap a shopper would accept (same use, "
                        "clearly better on what the rationale says), 1 = not a sensible swap. Leave blank if there is "
                        "no alternative. Rate from the product's real use, not from the system's score."),
    ("Afterwards", "Save as CSV with the same headers and run: python -m evaluation.pack_check "
                   "data/verified/pack_check/pack_check_sheet.csv"),
]


def get(base: str, path: str):
    with urllib.request.urlopen(base.rstrip("/") + path, timeout=30) as r:
        return json.load(r)


def fmt(v):
    if v is None:
        return ""
    return f"{v:g}" if isinstance(v, float) else v


def row_for(base: str, prod: dict) -> dict:
    d = get(base, f"/api/products/{prod['product_id']}")
    rep = d["report"]
    p = rep["product"]
    nut = p["nutrition"]
    items = (rep.get("ingredients_normalised") or {}).get("items") or []
    claims = ((rep.get("dimensions") or {}).get("claims") or {}).get("claims") or []
    alts = get(base, f"/api/products/{prod['product_id']}/alternatives?limit=3").get("alternatives") or []
    code = prod["source_url"].rstrip("/").split("/")[-1]
    r = {"category": prod["category"], "product_id": prod["product_id"], "off_code": code,
         "off_url": prod["source_url"], "name": prod["name"], "brand": prod["brand"] or "",
         "pack_size_stored": f"{fmt(prod['pack_size'])} {prod['pack_unit'] or ''}".strip(),
         "sys_basis": p.get("nutrition_basis") or "",
         "sys_ingredients": (p.get("ingredients_raw") or {}).get("value") or "",
         "sys_ingredient_count": len(items) if items else "",
         "sys_ingredients_uncertain": "; ".join(i.get("text") or i.get("name") for i in items
                                                if i.get("status") == "uncertain"),
         "sys_claims": "; ".join(f"{c.get('claim_text') or c.get('claim')}: {c.get('status')}" for c in claims),
         "sys_overall": fmt(rep["overall"].get("score")), "sys_grade": rep["overall"].get("grade") or "",
         "sys_data_status": d["status"]}
    for n, u in NUTRIENTS:
        r[f"sys_{n}_{u}"] = fmt((nut.get(n) or {}).get("value"))
    for i in range(3):
        a = alts[i] if i < len(alts) else None
        r[f"sys_alt_{i + 1}"] = f"{a['name']} ({a['brand']}), grade {a['grade']}" if a else ""
    return r


def sample(base: str) -> list[dict]:
    rng = random.Random(SEED)
    rows = []
    for cat, quota in QUOTA.items():
        prods = [p for p in get(base, f"/api/products?category={cat}&limit=500")["products"] if not p["is_synthetic"]]
        prods.sort(key=lambda p: p["product_id"])
        picked = rng.sample(prods, min(len(prods), quota + RESERVE))
        for k, prod in enumerate(picked):
            r = row_for(base, prod)
            r["sample"] = "main" if k < quota else "reserve"
            r["reserve_order"] = "" if k < quota else k - quota + 1
            rows.append(r)
    rows.sort(key=lambda r: (r["sample"] != "main", list(QUOTA).index(r["category"]), r["reserve_order"] or 0))
    for i, r in enumerate(rows, 1):
        r["check_id"] = f"PC-{i:03d}"
    return rows


def write_csv(rows: list[dict], path: Path) -> None:
    cols = ID_COLS + SYS_COLS + CHECK_COLS
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Check"
    cols = ID_COLS + SYS_COLS + CHECK_COLS
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    fills = {"id": "E8E8E8", "sys": "DCE9F7", "check": "FFF4CC"}
    for j, c in enumerate(cols, 1):
        kind = "id" if c in ID_COLS else "sys" if c in SYS_COLS else "check"
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fills[kind])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        width = 48 if c in ("sys_ingredients", "ingredients_notes", "sys_claims", "label_claims", "notes",
                            "sys_ingredients_uncertain", "rec_notes") else 30 if c.startswith("sys_alt") or c in (
                            "name", "off_url") else 14
        ws.column_dimensions[get_column_letter(j)].width = width
        if kind == "check":
            for i in range(2, len(rows) + 2):
                ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=fills["check"])
        if c in CHOICES:
            dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES[c]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{len(rows) + 1}")
        elif c.startswith("label_") and c not in ("label_claims", "label_pack_size", "label_basis"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{len(rows) + 1}")
    ws.freeze_panes = "J2"
    ws.auto_filter.ref = ws.dimensions

    g = wb.create_sheet("How to check", 0)
    g.column_dimensions["A"].width = 18
    g.column_dimensions["B"].width = 110
    g.append(["Pack-checking sheet", f"Generated {date.today().isoformat()} from methodology v0.1.0"])
    g["A1"].font = Font(bold=True, size=13)
    g.append([])
    for k, v in GUIDE:
        g.append([k, v])
        g.cell(row=g.max_row, column=1).font = Font(bold=True)
        g.cell(row=g.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    g.append([])
    g.append(["Column colours", "Grey: identification. Blue: what the system stored (do not edit). Yellow: fill from the pack."])
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--api", default="http://localhost:8000")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    rows = sample(args.api)
    write_csv(rows, OUT / "pack_check_sheet.csv")
    try:
        write_xlsx(rows, OUT / "pack_check_sheet.xlsx")
    except ImportError:
        print("openpyxl not installed; wrote CSV only")
    main_n = sum(r["sample"] == "main" for r in rows)
    print(f"{main_n} main + {len(rows) - main_n} reserve rows -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
