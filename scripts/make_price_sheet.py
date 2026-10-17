"""Build the price-entry sheet: one Zepto row and one Blinkit row for each of the 150 pack-check products.

    python scripts/make_price_sheet.py      # reads data/verified/pack_check/pack_check_sheet.csv

Prices are recorded by a person looking at the app or the shelf, never scraped. Extra rows can be added for
any product in the database (by off_code), for example the MRP printed on a pack.
"""
from __future__ import annotations

import csv
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PACK = ROOT / "data" / "verified" / "pack_check" / "pack_check_sheet.csv"
OUT = ROOT / "data" / "verified" / "prices"

ID_COLS = ["price_id", "check_id", "category", "off_code", "name", "brand", "pack_size_stored"]
ENTRY_COLS = ["platform", "observed_at", "city_pincode", "listing_name", "listing_pack_size", "selling_price_inr",
              "mrp_inr", "in_stock", "matches_product", "listing_url", "evidence_ref", "recorded_by", "notes"]
CHOICES = {
    "platform": ["zepto", "blinkit", "instamart", "bigbasket", "shop", "pack mrp", "brand website", "other"],
    "in_stock": ["yes", "no"],
    "matches_product": ["exact", "different size", "different variant", "not listed"],
}
GUIDE = [
    ("What this is", "One row per product per place a price was seen. Two rows (Zepto, Blinkit) are pre-made for "
                     "each of the 150 pack-check products; add rows for other platforms, a shop shelf, or the MRP "
                     "printed on a pack (platform = pack mrp)."),
    ("How to record", "Open the app as a normal shopper, search the product, and copy what is shown: the selling "
                      "price you would pay, the MRP if shown, and the pack size. Do not use scripts, bots or "
                      "browser extensions to collect prices, and do not copy prices in bulk."),
    ("Always fill", "observed_at (YYYY-MM-DD HH:MM), city_pincode (prices differ by area), listing_pack_size "
                    "(e.g. 200 g, 1 L, 4 x 70 g), selling_price_inr and matches_product. A row missing any of "
                    "these is not imported."),
    ("Same product?", "exact = same product and size; different size = same product, other pack size (still "
                      "usable, the price per 100 g is computed from the listed size); different variant or not "
                      "listed = not imported. Never enter a price for a product you did not see."),
    ("Evidence", "Keep a screenshot for each price, named by price_id, in the team drive, and write its name in "
                 "evidence_ref. Screenshots stay out of git."),
    ("Afterwards", "Save as CSV and run: PYTHONPATH=src python -m database.prices "
                   "data/verified/prices/price_entry_sheet.csv (imports, then rescores value for money)."),
]


def main() -> None:
    with PACK.open(encoding="utf-8") as f:
        products = [r for r in csv.DictReader(f) if r["sample"] == "main"]
    rows = []
    for p in products:
        for platform in ("zepto", "blinkit"):
            rows.append({c: p.get(c, "") for c in ID_COLS[1:]} | {"platform": platform})
    for i, r in enumerate(rows, 1):
        r["price_id"] = f"PR-{i:03d}"
    OUT.mkdir(parents=True, exist_ok=True)
    cols = ID_COLS + ENTRY_COLS
    with (OUT / "price_entry_sheet.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows({c: r.get(c, "") for c in cols} for r in rows)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    g = wb.active
    g.title = "How to record"
    g.column_dimensions["A"].width = 18
    g.column_dimensions["B"].width = 110
    g.append(["Price-entry sheet", "Prices recorded by hand for a college research project; nothing is scraped"])
    g["A1"].font = Font(bold=True, size=13)
    g.append([])
    for k, v in GUIDE:
        g.append([k, v])
        g.cell(row=g.max_row, column=1).font = Font(bold=True)
        g.cell(row=g.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws = wb.create_sheet("Prices")
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    n = len(rows) + 400  # room for added rows
    for j, c in enumerate(cols, 1):
        letter = get_column_letter(j)
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8E8E8" if c in ID_COLS else "FFF4CC")
        ws.column_dimensions[letter].width = 34 if c in ("name", "listing_name", "listing_url", "notes") else 15
        if c in CHOICES:
            dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES[c]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{n}")
        if c in ("selling_price_inr", "mrp_inr"):
            dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{n}")
    ws.freeze_panes = "H2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    wb.save(OUT / "price_entry_sheet.xlsx")
    print(f"{len(rows)} rows for {len(products)} products -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
